"""Layer 3: export the star schema for Power BI.

Power BI on the web cannot read a DuckDB file, but it can import an Excel
workbook. Each table goes on its own sheet as a named Excel table, so Power
BI picks every one up as a table of the semantic model. CSVs are written as
well for Power BI Desktop or any other tool.
"""
import duckdb
import pandas as pd
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from opslens.settings import POWERBI_EXPORT_DIR

TABLES = {
    "DimCustomer": "SELECT * EXCLUDE (has_telecoms, in_crm) FROM core.dim_customer",
    "DimProduct": "SELECT * FROM core.dim_product",
    "DimMonth": "SELECT * FROM core.dim_month",
    "FactMRR": "SELECT month, customer_ref, product_code, service_line, quantity, mrr::DOUBLE AS mrr FROM core.fct_mrr_month",
    "FactInvoice": """SELECT invoice_no, invoice_date, month, customer_ref, product_code, service_line, recurring,
                             line_total_gbp::DOUBLE AS amount_gbp, due_date, paid_date, days_to_pay, is_overdue
                      FROM core.fct_invoice_line""",
    "FactTicket": """SELECT ticket_id, customer_ref, month, opened_at, resolved_at, priority, category, status,
                            target_hours, resolution_hours, sla_status, csat_score, first_response_minutes
                     FROM core.fct_ticket""",
    "FactTelecomUsage": "SELECT * FROM core.fct_telecom_usage_month",
    "CustomerHealth": "SELECT * EXCLUDE (customer_name, sector, town, account_manager, customer_since) FROM mart.customer_health",
    "MRRMovements": "SELECT * FROM mart.mrr_monthly",
    "Tasks": "SELECT task_id, rule_id, title, owner, priority, status, created_at FROM ops.tasks",
}


def export(con: duckdb.DuckDBPyConnection) -> dict:
    POWERBI_EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    counts = {}
    for name, sql in TABLES.items():
        df = con.execute(sql).df()
        for col in df.select_dtypes(include=["datetimetz"]).columns:
            df[col] = df[col].dt.tz_localize(None)
        df.to_csv(POWERBI_EXPORT_DIR / f"{name}.csv", index=False)
        ws = wb.create_sheet(name)
        ws.append(list(df.columns))
        for row in df.itertuples(index=False):
            ws.append([None if pd.isna(v) else (v.to_pydatetime() if isinstance(v, pd.Timestamp) else v) for v in row])
        ref = f"A1:{get_column_letter(len(df.columns))}{max(len(df), 1) + 1}"
        table = Table(displayName=name, ref=ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
        counts[name] = len(df)
    wb.save(POWERBI_EXPORT_DIR / "OpsLens_PowerBI.xlsx")
    return counts
